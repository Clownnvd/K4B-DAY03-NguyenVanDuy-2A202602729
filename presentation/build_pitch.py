"""Build a concise FiProve-inspired 16:9 deck for the K4B Lab 3 pitch."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).with_name("Day03_K4B_Kingpro_Pitch.pptx")

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(23, 27, 35)
MUTED = RGBColor(97, 111, 129)
BLUE = RGBColor(49, 104, 231)
LIGHT_BLUE = RGBColor(241, 246, 255)
PALE = RGBColor(248, 249, 251)
BORDER = RGBColor(223, 228, 236)
GREEN = RGBColor(22, 128, 91)
LIGHT_GREEN = RGBColor(238, 248, 243)
AMBER = RGBColor(173, 95, 30)
LIGHT_AMBER = RGBColor(255, 246, 233)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def shape(slide, x, y, w, h, fill=WHITE, border=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    if border:
        item.line.color.rgb = border
        item.line.width = Pt(1)
    else:
        item.line.fill.background()
    if radius and item.adjustments:
        item.adjustments[0] = 0.14
    return item


def text(slide, x, y, w, h, value, size=17, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.text = line
    return box


def line(slide, x1, y1, x2, y2, color=BORDER, width=1.5):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def base(slide, number, section):
    shape(slide, 0, 0, 13.333, 7.5, WHITE, radius=False)
    shape(slide, 0, 0, 2.18, 7.5, PALE, radius=False)
    line(slide, 2.18, 0, 2.18, 7.5)
    shape(slide, 0.27, 0.26, 0.44, 0.44, BLUE)
    text(slide, 0.38, 0.345, 0.22, 0.25, "K", 15, WHITE, True,
         align=PP_ALIGN.CENTER)
    text(slide, 0.78, 0.28, 1.33, 0.28, "KINGPRO", 15, INK, True)
    text(slide, 0.28, 0.82, 1.72, 0.55, "HỌC VỤ · REACT\nLAB 03 / K4B", 9, MUTED, True)
    text(slide, 0.28, 2.0, 1.55, 0.22, "NỘI DUNG", 9, MUTED, True)
    labels = ["01  Bài toán", "02  Agentic Fit", "03  Kiến trúc",
              "04  Công cụ", "05  Demo", "06  Q&A"]
    for i, label in enumerate(labels, 1):
        yy = 2.33 + (i - 1) * 0.45
        if i == number:
            shape(slide, 0.18, yy - 0.04, 1.83, 0.37, LIGHT_BLUE)
        text(slide, 0.32, yy, 1.64, 0.25, label, 11,
             BLUE if i == number else MUTED, i == number)
    line(slide, 0.18, 6.84, 1.99, 6.84)
    shape(slide, 0.28, 6.95, 0.35, 0.35, RGBColor(232, 93, 77))
    text(slide, 0.34, 7.045, 0.23, 0.14, "ND", 8, WHITE, True,
         align=PP_ALIGN.CENTER)
    text(slide, 0.74, 6.96, 1.36, 0.19, "NGUYỄN VĂN DUY", 8, INK, True)
    text(slide, 0.74, 7.14, 1.2, 0.17, "2A202602729", 8, MUTED)
    text(slide, 2.55, 0.34, 5.8, 0.2, section.upper(), 10, BLUE, True)
    text(slide, 12.12, 7.07, 0.55, 0.2, f"{number:02d} / 06",
         9, MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=BLUE, fill=WHITE):
    shape(slide, x, y, w, h, fill, BORDER)
    shape(slide, x + 0.22, y + 0.23, 0.09, 0.38, accent)
    text(slide, x + 0.45, y + 0.21, w - 0.7, 0.37, title, 18, INK, True)
    text(slide, x + 0.25, y + 0.78, w - 0.5, h - 0.95, body, 14, MUTED)


# 01 — Problem and before/after
s = prs.slides.add_slide(blank)
base(s, 1, "Bài toán")
text(s, 2.55, 1.05, 9.9, 1.25, "Một câu hỏi.\nHai cách xử lý.", 37, INK, True)
text(s, 2.58, 2.65, 9.6, 0.62,
     "Câu hỏi chung cần câu trả lời. Yêu cầu tra cứu rồi tạo lịch cần dữ liệu,\nquyết định theo observation và một hành động có điều kiện.", 18, MUTED)
card(s, 2.58, 3.63, 4.55, 2.22, "Chatbot baseline",
     "Một lượt LLM · không tool.\nGiải thích chức năng demo; không xác minh được GPA hay cố vấn.", RGBColor(148, 160, 177))
card(s, 7.38, 3.63, 5.12, 2.22, "ReAct agent",
     "Tra cứu hồ sơ → đọc cố vấn → tạo bản nháp lịch.\nChỉ đi bước sau khi tool trả kết quả hợp lệ.", BLUE, LIGHT_BLUE)
shape(s, 2.58, 6.2, 9.92, 0.46, LIGHT_GREEN)
text(s, 2.88, 6.325, 9.25, 0.2,
     "DỮ LIỆU SINH VIÊN & LỊCH HẸN LÀ DEMO · OPENAI API DÙNG TRONG TEST LÀ THẬT",
     10, GREEN, True)

# 02 — Agentic Fit
s = prs.slides.add_slide(blank)
base(s, 2, "Agentic Fit")
text(s, 2.55, 0.72, 9.7, 0.6, "Bài này có cần agent không?", 30, INK, True)
text(s, 2.55, 1.35, 9.7, 0.4,
     "Chấm 4 tiêu chí trên thang 1–5 theo báo cáo Lab 3.", 16, MUTED)
fit = [
    ("Multi-step Reasoning", "4/5", "Tra cứu trước; dùng tên cố vấn từ kết quả để tạo lịch."),
    ("Tool Interaction", "5/5", "Một tool đọc hồ sơ, một tool tạo bản nháp."),
    ("Dynamic Decision", "4/5", "NOT_FOUND thì dừng; SUCCESS và đủ giờ mới đi tiếp."),
    ("Long Horizon Goal", "2/5", "Giữ mục tiêu vài vòng; chưa có memory dài hạn."),
]
positions = [(2.58, 2.05), (7.56, 2.05), (2.58, 3.75), (7.56, 3.75)]
for (title, score, reason), (x, y) in zip(fit, positions):
    shape(s, x, y, 4.72, 1.47, WHITE, BORDER)
    text(s, x + 0.22, y + 0.21, 3.45, 0.28, title, 16, INK, True)
    text(s, x + 3.65, y + 0.17, 0.78, 0.42, score, 21, BLUE, True,
         align=PP_ALIGN.RIGHT)
    text(s, x + 0.22, y + 0.72, 4.22, 0.5, reason, 12, MUTED)
shape(s, 2.58, 5.65, 9.7, 0.75, LIGHT_BLUE)
text(s, 2.86, 5.85, 2.6, 0.39, "TỔNG: 15/20", 21, BLUE, True)
text(s, 5.42, 5.9, 6.58, 0.35,
     "Theo bảng 3 tiêu chí trên slide: 13/15 → agent đáng thử.", 14, INK)
text(s, 2.6, 6.64, 9.45, 0.28,
     "FAQ nội bộ HR trong slide chỉ 3/15; dùng chatbot/retrieval là hợp lý hơn.", 11, MUTED)

# 03 — Architecture
s = prs.slides.add_slide(blank)
base(s, 3, "Kiến trúc")
text(s, 2.55, 0.72, 10.0, 0.55, "Agent nhìn kết quả rồi mới hành động", 28, INK, True)
text(s, 2.55, 1.31, 9.8, 0.4,
     "TC04 có chuỗi tool thật: academic_query → schedule_appointment → final answer.", 15, MUTED)
nodes = [(2.62, "Câu hỏi", PALE), (4.73, "LLM chọn bước", LIGHT_BLUE),
         (7.16, "academic_query", LIGHT_BLUE), (9.67, "Observation", LIGHT_GREEN)]
widths = [1.63, 1.89, 2.03, 2.15]
for (x, label, fill), w in zip(nodes, widths):
    shape(s, x, 2.21, w, 0.78, fill, BORDER)
    text(s, x + 0.1, 2.45, w - 0.2, 0.26, label, 14, INK, True,
         align=PP_ALIGN.CENTER)
for x in [4.34, 6.74, 9.29]:
    text(s, x, 2.42, 0.29, 0.32, "→", 19, BLUE, True)
shape(s, 2.63, 3.58, 4.19, 1.43, LIGHT_AMBER, BORDER)
text(s, 2.9, 3.83, 3.65, 0.25, "NOT_FOUND / tool lỗi", 16, AMBER, True)
text(s, 2.9, 4.24, 3.65, 0.49,
     "Dừng; nói rõ không có hồ sơ.\nKhông tạo lịch, không bịa cố vấn.", 12, MUTED)
shape(s, 7.05, 3.58, 5.49, 1.43, LIGHT_GREEN, BORDER)
text(s, 7.32, 3.82, 4.95, 0.27, "SUCCESS + người dùng yêu cầu đặt lịch", 15, GREEN, True)
text(s, 7.32, 4.22, 4.91, 0.56,
     "Lấy đúng cố vấn → schedule_appointment\n→ trả lời: bản nháp DEMO, chưa gửi VinUni.", 12, MUTED)
text(s, 6.47, 3.14, 0.38, 0.28, "↙", 18, AMBER, True)
text(s, 10.23, 3.14, 0.38, 0.28, "↘", 18, GREEN, True)
shape(s, 2.62, 5.43, 9.92, 0.84, PALE, BORDER)
text(s, 2.9, 5.62, 9.36, 0.42,
     "Baseline: Câu hỏi → một lượt LLM → câu trả lời. Không có tool hoặc observation.", 14, INK)
text(s, 2.63, 6.52, 9.91, 0.45,
     "MCP-style adapter nằm trong cùng tiến trình theo starter repo; giới hạn 5 vòng và chặn tool call lặp.", 11, MUTED)

# 04 — Tools and guardrails
s = prs.slides.add_slide(blank)
base(s, 4, "Công cụ")
text(s, 2.55, 0.72, 9.9, 0.55, "Hai công cụ, hai trách nhiệm rõ ràng", 28, INK, True)
text(s, 2.55, 1.35, 9.8, 0.35,
     "Model chọn tool; backend quyết định có cho phép hành động hay không.", 15, MUTED)
shape(s, 2.58, 1.95, 4.7, 3.75, WHITE, BORDER)
shape(s, 2.83, 2.18, 0.44, 0.44, LIGHT_BLUE)
text(s, 2.96, 2.29, 0.19, 0.18, "1", 12, BLUE, True, align=PP_ALIGN.CENTER)
text(s, 3.44, 2.22, 3.42, 0.38, "academic_query", 19, INK, True)
text(s, 2.86, 2.88, 4.08, 0.72,
     "Đầu vào: student_id.\nĐọc hai hồ sơ DEMO theo mã.", 15, MUTED)
line(s, 2.85, 3.91, 6.95, 3.91)
text(s, 2.86, 4.18, 4.08, 1.16,
     "Đầu ra: SUCCESS kèm GPA/cố vấn, hoặc NOT_FOUND. Không lấy dữ liệu VinUni thật.", 14, INK)
shape(s, 7.52, 1.95, 5.02, 3.75, LIGHT_BLUE, BORDER)
shape(s, 7.79, 2.18, 0.44, 0.44, WHITE)
text(s, 7.92, 2.29, 0.19, 0.18, "2", 12, BLUE, True, align=PP_ALIGN.CENTER)
text(s, 8.43, 2.22, 3.78, 0.38, "schedule_appointment", 18, INK, True)
text(s, 7.8, 2.88, 4.36, 0.7,
     "Đầu vào: mã, giờ hẹn, tên cố vấn.\nTạo bản nháp lịch trong bộ nhớ.", 15, MUTED)
line(s, 7.81, 3.91, 12.22, 3.91)
text(s, 7.8, 4.18, 4.4, 1.18,
     "Chỉ SUCCESS nếu đã tra cứu cùng phiên, đúng cố vấn và giờ đúng định dạng. Không đặt lịch thật.", 14, INK)
shape(s, 2.58, 6.04, 9.96, 0.62, LIGHT_GREEN)
text(s, 2.86, 6.2, 9.4, 0.29,
     "PHANH AN TOÀN: MAX 5 VÒNG · CHẶN TOOL/ARGS LẶP · LỖI THÌ DỪNG / GIẢI THÍCH",
     11, GREEN, True)

# 05 — Live demo
s = prs.slides.add_slide(blank)
base(s, 5, "Demo")
text(s, 2.55, 0.72, 10.0, 0.54, "Demo 2 câu và mở trace từng bước", 28, INK, True)
text(s, 2.55, 1.32, 9.75, 0.35,
     "Mở giao diện: http://127.0.0.1:8765 · bên trái chọn Bộ 5 tình huống.", 14, MUTED)
image_path = ROOT / "docs" / "ui_live.png"
if image_path.exists():
    s.shapes.add_picture(str(image_path), Inches(2.57), Inches(1.94), width=Inches(6.17))
shape(s, 8.94, 1.94, 3.57, 1.45, WHITE, BORDER)
text(s, 9.18, 2.14, 3.1, 0.28, "TC01 · CÂU HỎI CHUNG", 13, BLUE, True)
text(s, 9.18, 2.57, 3.1, 0.53,
     "Giải thích chức năng demo.\n0 tool call → chatbot đủ.", 13, MUTED)
shape(s, 8.94, 3.56, 3.57, 1.77, LIGHT_BLUE, BORDER)
text(s, 9.18, 3.77, 3.1, 0.28, "TC04 · YÊU CẦU ĐẶT LỊCH", 13, BLUE, True)
text(s, 9.18, 4.16, 3.1, 0.96,
     "Tra cứu SV2026002 → đúng cố vấn → bản nháp lịch.\n2 tool calls; mở trace bên phải.", 13, MUTED)
shape(s, 2.58, 6.16, 9.94, 0.59, LIGHT_GREEN)
text(s, 2.88, 6.33, 9.34, 0.24,
     "5/5 CA TEST API THẬT · 4 TOOL CALLS ĐÚNG · TC05 NOT_FOUND KHÔNG TẠO LỊCH",
     11, GREEN, True)

# 06 — Q&A backup
s = prs.slides.add_slide(blank)
base(s, 6, "Q&A")
text(s, 2.55, 0.72, 9.9, 0.56, "Những giới hạn cần nói thẳng", 28, INK, True)
qa = [
    ("Dữ liệu thật?", "Không. Hồ sơ và lịch hẹn là DEMO; chỉ OpenAI API dùng để nghiệm thu là thật."),
    ("MCP chạy qua mạng?", "Chưa. Đây là adapter MCP-style trong cùng tiến trình theo starter repo."),
    ("Thought có phải CoT?", "Không. Trace ghi quyết định gọi tool quan sát được, args, observation và final."),
    ("Agent luôn tốt hơn?", "Không. Câu hỏi đơn giản đi đường chatbot; agent có thêm cost/latency."),
]
for i, (question, answer) in enumerate(qa):
    y = 1.66 + i * 1.14
    shape(s, 2.58, y, 9.94, 0.94, WHITE if i % 2 == 0 else PALE, BORDER)
    text(s, 2.83, y + 0.17, 2.28, 0.4, question, 14, BLUE, True)
    text(s, 5.38, y + 0.15, 6.72, 0.6, answer, 13, INK)
shape(s, 2.58, 6.43, 9.94, 0.46, LIGHT_BLUE)
text(s, 2.87, 6.55, 9.35, 0.2,
     "Repo K4B đã nộp VLearn · 15/20 Agentic Fit · 5/5 test là kết quả test, không phải điểm thầy.",
     10, BLUE, True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"{OUT} | {len(prs.slides)} slides | {OUT.stat().st_size} bytes")
